import os
import json
from typing import List
from fastapi import APIRouter, HTTPException
from fastapi.responses import FileResponse
from pydantic import BaseModel
from pptx import Presentation

from db.queries import get_slide, update_shape_text, get_shapes_for_slide, get_slides_for_pptx
from db.connection import get_db
from parsers.slide_renderer import export_slide_to_png
from parsers.pptx_parser import serialize_slide

router = APIRouter()


class ShapeUpdate(BaseModel):
    id: str
    full_text: str


@router.get("/")
async def get_slides(pptx_id: str, slide_index: int):
    slides_data = get_slides_for_pptx(pptx_id)
    if not slides_data:
        return {"message": "No slides found for the specified PPTX file"}
    for slide in slides_data:
        if slide["slide_index"] == slide_index:
            return {"slide": slide}
    return {"message": "Slide not found"}


@router.get("/{slide_id}/png")
async def get_slide_png(slide_id: str):
    """Serve the slide PNG from slide_id."""
    slide_data = get_slide(slide_id)
    if not slide_data:
        raise HTTPException(status_code=404, detail="Slide not found")

    png_path = slide_data.get("png_path")
    if not png_path or not os.path.exists(png_path):
        raise HTTPException(status_code=404, detail="Slide PNG image not found")

    return FileResponse(png_path, media_type="image/png")


@router.post("/{slide_id}/shapes")
async def update_slide_shapes(slide_id: str, updates: List[ShapeUpdate]):
    """
    Accepts updated texts for specific database shape IDs,
    writes them to the pptx file via python-pptx,
    re-exports the PNG slide, updates DB shapes, and returns the slide shapes.
    """
    slide_data = get_slide(slide_id)
    if not slide_data:
        raise HTTPException(status_code=404, detail="Slide not found")

    pptx_file_id = slide_data["pptx_file_id"]
    slide_index = slide_data["slide_index"]

    # Get PPTX storage path
    with get_db() as db:
        row = db.execute("SELECT storage_path FROM pptx_files WHERE id=?", (pptx_file_id,)).fetchone()
        pptx_path = row["storage_path"] if row else None

    if not pptx_path or not os.path.exists(pptx_path):
        raise HTTPException(status_code=404, detail="Original PPTX file not found")

    try:
        # Load and modify PPTX
        prs = Presentation(pptx_path)
        slide_obj = prs.slides[slide_index]

        for update in updates:
            # Get pptx_shape_id from DB
            with get_db() as db:
                srow = db.execute("SELECT pptx_shape_id FROM shapes WHERE id=?", (update.id,)).fetchone()
                pptx_shape_id = srow["pptx_shape_id"] if srow else None

            if pptx_shape_id is None:
                continue

            # Find matching pptx shape
            target_shape = None
            for sh in slide_obj.shapes:
                if sh.shape_id == pptx_shape_id:
                    target_shape = sh
                    break

            if target_shape and target_shape.has_text_frame:
                bold = False
                italic = False
                font_size = None
                color_rgb = None
                align = None
                try:
                    if target_shape.text_frame.paragraphs and target_shape.text_frame.paragraphs[0].runs:
                        first_run = target_shape.text_frame.paragraphs[0].runs[0]
                        bold = first_run.font.bold
                        italic = first_run.font.italic
                        font_size = first_run.font.size
                        color_rgb = first_run.font.color.rgb
                    if target_shape.text_frame.paragraphs:
                        align = target_shape.text_frame.paragraphs[0].alignment
                except Exception:
                    pass

                # Write new text
                target_shape.text_frame.text = update.full_text

                # Restore style on the new run
                try:
                    if target_shape.text_frame.paragraphs and target_shape.text_frame.paragraphs[0].runs:
                        new_run = target_shape.text_frame.paragraphs[0].runs[0]
                        new_run.font.bold = bold
                        new_run.font.italic = italic
                        if font_size:
                            new_run.font.size = font_size
                        if color_rgb:
                            new_run.font.color.rgb = color_rgb
                    if align is not None and target_shape.text_frame.paragraphs:
                        target_shape.text_frame.paragraphs[0].alignment = align
                except Exception:
                    pass

        # Save presentation
        prs.save(pptx_path)

        # Re-render PNG
        png_path = slide_data.get("png_path")
        if not png_path:
            processed_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "processed"))
            slides_dir = os.path.join(processed_dir, "slides")
            os.makedirs(slides_dir, exist_ok=True)
            png_path = os.path.join(slides_dir, f"{slide_id}.png")
            with get_db() as db:
                db.execute("UPDATE slides SET png_path=? WHERE id=?", (png_path, slide_id))

        export_slide_to_png(pptx_path, slide_index, png_path)

        # Re-serialize shapes and update DB
        updated_shapes = serialize_slide(slide_obj, prs.slide_width, prs.slide_height)
        for ush in updated_shapes:
            with get_db() as db:
                db_row = db.execute(
                    "SELECT id FROM shapes WHERE slide_id=? AND pptx_shape_id=?",
                    (slide_id, ush["pptx_shape_id"])
                ).fetchone()
                if db_row:
                    db_shape_id = db_row["id"]
                    update_shape_text(db_shape_id, ush["full_text"], ush["runs_json"])

        # Fetch and return updated shapes
        shapes = get_shapes_for_slide(slide_id)
        for sh in shapes:
            sh["runs_json"] = json.loads(sh["runs_json"])

        return {
            "slide_id": slide_id,
            "slide_index": slide_index,
            "png_url": f"/api/slides/{slide_id}/png",
            "shapes": shapes
        }

    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

